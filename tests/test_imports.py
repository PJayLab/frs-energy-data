"""Import regression tests; SQLite exercises real async SQLAlchemy sessions.

PostGIS/JSONB storage is substituted only in this test process.
"""
import unittest
import os
import uuid
import asyncio
from io import BytesIO
from unittest.mock import patch

from sqlalchemy import JSON, Text, select, text
from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine
from sqlalchemy.types import TypeDecorator

from app.models import Object, ServiceConnection, ConnectionIssueReport
from app.schemas import GPSImportData, ImportData
from app import services
from app.api.import_router import import_excel, import_response
from app.import_helpers import ObjectIndex, find_reference_object, upsert_object
from fastapi import UploadFile, HTTPException
from openpyxl import Workbook


class GeometryText(TypeDecorator):
    impl = Text
    cache_ok = True

    def process_bind_param(self, value, dialect):
        return str(value) if value is not None else None


def row(objects='Chäppelimatt 5', notes='PVA: 30 kVA', municipality='A'):
    return (municipality, objects, None, notes, None, None, None, 'Transformer A', '4')


class ParsingTests(unittest.TestCase):
    def test_response_summarizes_warnings_without_changing_success_counts(self):
        report = dict(imported=[object(), object()], errors=[], warnings=[
            {'reason': 'reference_alias_matched'}, {'reason': 'reference_alias_matched'},
            {'reason': 'object_not_found'}], created=0, updated=1, unchanged=1, skipped=1)
        response = import_response(report)
        self.assertEqual(response['imported_count'], 2)
        self.assertEqual(response['warnings_by_reason'], {'reference_alias_matched': 2, 'object_not_found': 1})

    def test_single(self):
        self.assertEqual(services.normalize_row(row())[0]['connection_notes'], ['PVA: 30 kVA'])

    def test_equal_notes(self):
        entries = services.normalize_row(row('Haus 1\nHaus 2', 'Note 1\nNote 2'))
        self.assertEqual([e['connection_notes'] for e in entries], [['Note 1'], ['Note 2']])

    def test_positional_blanks(self):
        entries = services.normalize_row(row('Haus 1\nHaus 2\nHaus 3', '\n\nLadestation: 11 kVA'))
        self.assertEqual([e['connection_notes'] for e in entries], [[], [], ['Ladestation: 11 kVA']])

    def test_blank_object_and_insurance_positions(self):
        data = list(row('Haus 1\n\nHaus 3', 'First\n\nThird'))
        data[2] = '0333\n\n0335'
        entries = services.normalize_row(data)
        self.assertEqual([e['insurance_number'] for e in entries], ['0333', '0335'])
        self.assertEqual(entries[1]['connection_notes'], ['Third'])

    def test_ambiguous_overflow_is_not_broadcast_or_silently_truncated(self):
        technical = services.normalize_row(row('A\nB', 'PVA\nx\ny'))
        self.assertTrue(all(e['_skip_import'] for e in technical))
        ordinary = services.normalize_row(row('A\nB', 'x\ny\nz'))
        self.assertTrue(all(e['_skip_import'] for e in ordinary))
        self.assertTrue(ordinary[0]['_warnings'])

    def test_continuation_notes_and_leading_zeroes(self):
        data = list(row('Büntenweg 8\n\nBüntenweg 10', 'PVA: 12.5 kVA\nLadestation: 11 kVA'))
        data[2] = '0290\n\n0289'
        entries = services.normalize_row(data)
        self.assertEqual([e['insurance_numbers'] for e in entries], [['0290'], ['0289']])
        self.assertEqual([e['connection_notes'] for e in entries], [['PVA: 12.5 kVA', 'Ladestation: 11 kVA'], []])
        self.assertFalse(any(e['_skip_import'] for e in entries))

    def test_trailing_continuation_and_parallel_note(self):
        entries = services.normalize_row(row('A\nB\n', '\nPVA: 9 kVA\nLadestation: 11 kVA\nParallel Strang 1+2'))
        self.assertEqual(entries[0]['connection_notes'], ['Parallel Strang 1+2'])
        self.assertEqual(entries[1]['connection_notes'], ['PVA: 9 kVA', 'Ladestation: 11 kVA', 'Parallel Strang 1+2'])
        self.assertFalse(any(e['_skip_import'] for e in entries))

    def test_unassigned_number_is_not_shifted(self):
        for objects, numbers in [('A\n\nB', '01\n02\n03'), ('A', '01\n02')]:
            data = list(row(objects, None))
            data[2] = numbers
            entries = services.normalize_row(data)
            self.assertTrue(all(e['_skip_import'] for e in entries))
            self.assertEqual(entries[0]['_warnings'][0]['reason'], 'unassigned_insurance_number')

    def test_repeated_address_collects_numbers_and_notes(self):
        data = list(row('Feldmatt\n\nFeldmatt', 'Scheune\nPVA: 25 kVA'))
        data[2] = '0366A\n\n0366'
        entries = services.normalize_row(data)
        self.assertEqual(len(entries), 1)
        self.assertEqual(entries[0]['insurance_numbers'], ['0366A', '0366'])
        self.assertEqual(entries[0]['connection_notes'], ['Scheune', 'PVA: 25 kVA'])

    def test_missing_number_remains_missing(self):
        data = list(row('A\nB\nC', None))
        data[2] = '01\n\n03'
        self.assertEqual([e['insurance_numbers'] for e in services.normalize_row(data)], [['01'], [], ['03']])

    def test_zero_crlf_nbsp(self):
        self.assertEqual(services.split_positional('\r\nx\xa0y\r\n'), ['', 'x y'])
        self.assertEqual(services.split_compact(0), ['0'])

    def test_multiple_source_edit_overrides_cached_names(self):
        self.assertEqual(services.source_values(dict(source_name='A\nB', source_names=['A'])), ['A', 'B'])


class ImportTests(unittest.IsolatedAsyncioTestCase):
    async def asyncSetUp(self):
        self.types = [(Object.__table__.c.geom, Object.__table__.c.geom.type)]
        Object.__table__.c.geom.type = GeometryText()
        column = Object.__table__.c.insurance_numbers
        self.types.append((column, column.type))
        column.type = JSON()
        for name in ('source_outgoing', 'disconnect_point_outgoing', 'connection_notes'):
            column = ServiceConnection.__table__.c[name]
            self.types.append((column, column.type))
            column.type = JSON()
        self.engine = create_async_engine('sqlite+aiosqlite:///:memory:')
        async with self.engine.begin() as conn:
            await conn.run_sync(lambda c: Object.__table__.create(c))
            await conn.run_sync(lambda c: ServiceConnection.__table__.create(c))
            await conn.run_sync(lambda c: ConnectionIssueReport.__table__.create(c))
        self.sessions = async_sessionmaker(self.engine, expire_on_commit=False)
        self.session = self.sessions()
        self.session.add(Object(name='TransformerA', friendly_name='Transformer A', type='transformer'))
        await self.session.commit()

    async def asyncTearDown(self):
        await self.session.close()
        await self.engine.dispose()
        for column, original in self.types:
            column.type = original

    async def objects(self):
        return list((await self.session.scalars(select(Object))).all())

    async def connections(self):
        return list((await self.session.scalars(select(ServiceConnection))).all())

    async def seed_wrong_municipality_connections(self):
        geu = Object(name='Bahnhofplatz6', location='Geuensee', type='building', insurance_numbers=['0103B', '0208'])
        sur = Object(name='Bahnhofplatz6', location='Sursee', type='building')
        geu_ts = Object(name='GEU-Allmend', type='transformer')
        sur_ts = Object(name='SUR-Fruchthof', type='transformer')
        box = Object(name='Hinterdorfstrasse16', type='distribution_box', location='Geuensee')
        # Same-name grid roles are distinct under typed reference matching.
        sur_box = Object(name='SUR-Fruchthof', type='distribution_box')
        self.session.add_all([geu, sur, geu_ts, sur_ts, box, sur_box])
        await self.session.flush()
        correct = ServiceConnection(building_id=geu.id, transformer_id=geu_ts.id, distribution_box_id=box.id,
            source_outgoing=['4'], disconnect_point_outgoing=['5'], connection_notes=[])
        stale = [ServiceConnection(building_id=geu.id, transformer_id=sur_ts.id, distribution_box_id=sur_ts.id,
            source_outgoing=['2'], disconnect_point_outgoing=['2'], connection_notes=notes)
            for notes in (['Stationsgebäude'], ['Verkaufsstand'], [])]
        self.session.add_all([correct, *stale])
        await self.session.commit()
        entries = services.normalize_row(('Geuensee', 'Bahnhofplatz 6', '0103B', None,
            None, 'Hinterdorfstrasse 16', '5', 'GEU-Allmend', '4'))
        entries += services.normalize_row(('Sursee', 'Bahnhofplatz 6\nBahnhofplatz 6\nBahnhofplatz 6',
            '0208\n9886\n0208C', 'Stationsgebäude\nVerkaufsstand', None, 'SUR-Fruchthof', '2', 'SUR-Fruchthof', '2'))
        return geu.id, sur.id, correct.id, [c.id for c in stale], entries

    async def test_replace_repairs_bahnhofplatz_and_is_idempotent(self):
        geu_id, sur_id, correct_id, stale_ids, entries = await self.seed_wrong_municipality_connections()
        result = await services.import_service_connections(entries, self.session, replace_existing=True)
        self.assertEqual(result['deleted'], 3)
        self.assertEqual({c['id'] for c in result['removed_connections']}, set(stale_ids))
        connections = await self.connections()
        self.assertEqual(len(connections), 2)
        geu = next(c for c in connections if c.building_id == geu_id)
        sur = next(c for c in connections if c.building_id == sur_id)
        self.assertEqual(geu.id, correct_id)
        self.assertEqual(geu.connection_notes, [])
        self.assertEqual(sur.connection_notes, ['Stationsgebäude', 'Verkaufsstand'])
        self.assertEqual((await self.session.get(Object, geu_id)).insurance_numbers, ['0103B'])
        self.assertEqual((await self.session.get(Object, sur_id)).insurance_numbers, ['0208', '9886', '0208C'])
        ids = {c.id for c in connections}
        again = await services.import_service_connections(entries, self.session, replace_existing=True)
        self.assertEqual((again['created'], again['updated'], again['deleted']), (0, 0, 0))
        self.assertEqual({c.id for c in await self.connections()}, ids)

    async def test_replace_preview_rolls_back_everything(self):
        geu_id, sur_id, correct_id, stale_ids, entries = await self.seed_wrong_municipality_connections()
        result = await services.import_service_connections(entries, self.session, replace_existing=True, dry_run=True)
        self.assertEqual(result['deleted'], 3)
        self.assertEqual(import_response(result)['status'], 'preview')
        self.assertEqual({c.id for c in await self.connections()}, {correct_id, *stale_ids})
        self.assertEqual((await self.session.get(Object, geu_id)).insurance_numbers, ['0103B', '0208'])
        self.assertIsNone((await self.session.get(Object, sur_id)).insurance_numbers)

    async def test_replace_unresolved_row_protects_entire_building(self):
        geu_id, _, correct_id, stale_ids, entries = await self.seed_wrong_municipality_connections()
        bad = {**entries[0], 'source_name': 'Missing', 'source_names': []}
        result = await services.import_service_connections([entries[0], bad], self.session, replace_existing=True)
        self.assertEqual(result['deleted'], 0)
        self.assertEqual(result['skipped'], 2)
        self.assertEqual({c.id for c in await self.connections()}, {correct_id, *stale_ids})
        self.assertEqual((await self.session.get(Object, geu_id)).insurance_numbers, ['0103B', '0208'])

    async def test_replace_does_not_touch_other_municipality_or_unlisted_buildings(self):
        _, sur_id, correct_id, _, entries = await self.seed_wrong_municipality_connections()
        await services.import_service_connections(entries, self.session, replace_existing=True)
        before = next(c for c in await self.connections() if c.building_id == sur_id)
        sur_connection_id = before.id
        await services.import_service_connections([entries[0]], self.session, replace_existing=True)
        self.assertEqual({c.id for c in await self.connections()}, {correct_id, sur_connection_id})

    async def test_replace_issue_report_blocks_stale_connection_deletion(self):
        _, _, correct_id, stale_ids, entries = await self.seed_wrong_municipality_connections()
        issue = ConnectionIssueReport(connection_id=stale_ids[0], user='test', remarks='Review')
        self.session.add(issue)
        await self.session.commit()
        result = await services.import_service_connections([entries[0]], self.session, replace_existing=True)
        self.assertEqual(result['deleted'], 0)
        self.assertEqual(result['warnings'][-1]['reason'], 'replacement_blocked_by_issue_reports')
        self.assertEqual({c.id for c in await self.connections()}, {correct_id, *stale_ids})

    async def test_replace_collapses_duplicates_and_combines_source_row_notes(self):
        entry = services.normalize_row(row())[0]
        await services.import_service_connections([entry], self.session)
        original = (await self.connections())[0]
        self.session.add(ServiceConnection(building_id=original.building_id, transformer_id=original.transformer_id,
            source_outgoing=['4'], disconnect_point_outgoing=[], connection_notes=['Old']))
        await self.session.commit()
        result = await services.import_service_connections([entry, {**entry, 'connection_notes': ['Additional']}],
            self.session, replace_existing=True)
        self.assertEqual(result['deleted'], 1)
        self.assertEqual(result['consolidated'], 1)
        self.assertEqual(len(await self.connections()), 1)
        self.assertEqual((await self.connections())[0].connection_notes, ['PVA: 30 kVA', 'Additional'])

    async def test_replace_clears_blank_optional_reference(self):
        geu_id, _, correct_id, _, entries = await self.seed_wrong_municipality_connections()
        entry = {**entries[0], 'first_disconnect_point': None}
        await services.import_service_connections([entry], self.session, replace_existing=True)
        connection = next(c for c in await self.connections() if c.building_id == geu_id)
        self.assertIsNone(connection.distribution_box_id)
        self.assertNotEqual(connection.id, correct_id)  # A changed path receives a new identity.

    async def test_replace_rejects_incomplete_payload_and_multiple_sources(self):
        entry = services.normalize_row(row())[0]
        for invalid in [{**entry, 'municipality': None}, {k: v for k, v in entry.items() if k != 'source_outgoing'}]:
            with self.assertRaises(ValueError):
                await services.import_service_connections([invalid], self.session, replace_existing=True)
        entry['source_name'] = 'Transformer A\nTransformer B'
        result = await services.import_service_connections([entry], self.session, replace_existing=True)
        self.assertEqual(result['skipped'], 1)
        self.assertEqual(await self.connections(), [])

    async def test_replace_failure_rolls_back_deletes(self):
        _, _, correct_id, stale_ids, entries = await self.seed_wrong_municipality_connections()
        with patch.object(self.session, 'commit', side_effect=RuntimeError('commit failed')):
            with self.assertRaises(RuntimeError):
                await services.import_service_connections(entries, self.session, replace_existing=True)
        self.assertEqual({c.id for c in await self.connections()}, {correct_id, *stale_ids})

    async def test_excel_replace_preview_then_apply(self):
        _, _, correct_id, stale_ids, _ = await self.seed_wrong_municipality_connections()
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = 'Modified'
        sheet.append(['Gemeinde', 'Objekt', 'Assek. Nr.', 'Bemerkungen/Verbindung',
            'TK ohne Schaltmöglichkeit', 'Erste Trennstelle', 'Abgang Trennstelle', 'Speisung', 'Abgang Speisung'])
        sheet.append(['Geuensee', 'Bahnhofplatz 6', '0103B', None, None, 'Hinterdorfstrasse 16', '5', 'GEU-Allmend', '4'])
        data = BytesIO()
        workbook.save(data)
        workbook.close()
        for preview in (True, False):
            result = await import_excel(UploadFile(file=BytesIO(data.getvalue()), filename='test.xlsx'),
                self.session, mode='replace_buildings', dry_run=preview)
            self.assertEqual(result['deleted'], 3)
            self.assertEqual(result['status'], 'preview' if preview else 'success')
            self.assertEqual({c.id for c in await self.connections()}, {correct_id, *stale_ids} if preview else {correct_id})

    async def test_excel_replace_rejects_unrecognized_columns(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = 'Modified'
        sheet.append(['Wrong layout'])
        data = BytesIO()
        workbook.save(data)
        workbook.close()
        with self.assertRaises(HTTPException) as raised:
            await import_excel(UploadFile(file=BytesIO(data.getvalue()), filename='test.xlsx'),
                self.session, mode='replace_buildings')
        self.assertEqual(raised.exception.status_code, 400)

    async def test_gps_repair_splits_building_and_is_repeatable(self):
        asset = Object(name='Schlottermilch10', location='A', type='building',
                       ckw_id='KAK10', insurance_numbers=['001'])
        self.session.add(asset)
        await self.session.flush()
        transformer = next(o for o in await self.objects() if o.type == 'transformer')
        connection = ServiceConnection(building_id=asset.id, transformer_id=transformer.id,
                                       distribution_box_id=asset.id)
        self.session.add(connection)
        await self.session.commit()
        original_id, connection_id = asset.id, connection.id
        payload = GPSImportData(points=[dict(name='Schlottermilch10', ckw_id='KAK10',
            type='distribution_box', lat=47, lon=8)])
        preview = await services.import_gps_objects(payload, self.session,
            return_report=True, repair_buildings=True, dry_run=True)
        self.assertEqual(len(preview['repairs']), 1)
        self.assertEqual(len(await self.objects()), 2)
        asset = await self.session.get(Object, original_id)
        self.assertEqual(asset.type, 'building')
        result = await services.import_gps_objects(payload, self.session,
            return_report=True, repair_buildings=True)
        self.assertEqual(result['updated'], 1)
        connection = await self.session.get(ServiceConnection, connection_id)
        building = await self.session.get(Object, connection.building_id)
        self.assertNotEqual(building.id, original_id)
        self.assertIsNone(building.geom)
        self.assertIsNone(building.ckw_id)
        self.assertEqual(building.insurance_numbers, ['001'])
        self.assertEqual(connection.distribution_box_id, original_id)
        self.assertEqual(asset.type, 'distribution_box')
        again = await services.import_gps_objects(payload, self.session,
            return_report=True, repair_buildings=True)
        self.assertEqual(again['repairs'], [])
        self.assertEqual(len(await self.objects()), 3)

    async def test_gps_then_excel_keeps_same_address_separate(self):
        payload = GPSImportData(points=[dict(name='Schlottermilch10', ckw_id='KAK10',
            type='distribution_box', lat=47, lon=8)])
        await services.import_gps_objects(payload, self.session)
        entry = services.normalize_row(row('Schlottermilch 10'))[0]
        entry['first_disconnect_point'] = 'Schlottermilch10'
        await services.import_service_connections([entry], self.session)
        await services.import_gps_objects(payload, self.session)
        await services.import_service_connections([entry], self.session)
        connection = (await self.connections())[0]
        self.assertNotEqual(connection.building_id, connection.distribution_box_id)
        self.assertEqual(len(await self.objects()), 3)
        self.assertEqual(len(await self.connections()), 1)
        building = await self.session.get(Object, connection.building_id)
        self.assertIsNone(building.geom)

    async def test_gps_repair_reuses_existing_building(self):
        asset = Object(name='Haus1', type='building', ckw_id='KAK1', location='A')
        building = Object(name='Haus1', type='building', location='A')
        self.session.add_all([asset, building])
        await self.session.commit()
        payload = GPSImportData(points=[dict(name='Haus1', ckw_id='KAK1',
            type='distribution_box', lat=47, lon=8)])
        normal = await services.import_gps_objects(payload, self.session, return_report=True)
        self.assertEqual(normal['skipped'], 1)
        result = await services.import_gps_objects(payload, self.session,
            return_report=True, repair_buildings=True)
        self.assertEqual(result['repairs'][0]['building_id'], building.id)
        self.assertEqual(len(await self.objects()), 3)

    async def test_reference_requires_correct_grid_type(self):
        self.session.add(Object(name='WrongBox', type='transformer', location='A'))
        await self.session.commit()
        entry = services.normalize_row(row())[0]
        entry['first_disconnect_point'] = 'WrongBox'
        result = await services.import_service_connections([entry], self.session)
        self.assertEqual(result['skipped'], 1)
        self.assertEqual(await self.connections(), [])

    async def test_building_and_grid_asset_with_same_name_remain_separate(self):
        box = Object(name='Haus1', location='A', type='distribution_box')
        self.session.add(box)
        await self.session.commit()
        entry = services.normalize_row(row('Haus 1'))[0]
        entry['first_disconnect_point'] = 'Haus 1'
        result = await services.import_service_connections([entry], self.session)
        self.assertEqual(result['created'], 1)
        connection = (await self.connections())[0]
        self.assertNotEqual(connection.building_id, box.id)
        self.assertEqual(connection.distribution_box_id, box.id)
        self.assertEqual(box.type, 'distribution_box')

    async def test_building_asset_id_cannot_move_between_municipalities(self):
        self.session.add(Object(name='Haus1', location='A', ckw_id='same', type='building'))
        await self.session.commit()
        entry = services.normalize_row(row('Haus 1', municipality='B'))[0]
        result = await services.import_service_connections([{**entry, 'ckw_id': 'same'}], self.session)
        self.assertEqual(result['skipped'], 1)
        self.assertEqual(next(o for o in await self.objects() if o.ckw_id == 'same').location, 'A')

    async def test_unknown_location_not_adopted_across_two_municipalities(self):
        historical = Object(name='Haus1', type='building')
        self.session.add(historical)
        await self.session.commit()
        entries = services.normalize_row(row('Haus 1', municipality='A')) + services.normalize_row(row('Haus 1', municipality='B'))
        await services.import_service_connections(entries, self.session)
        self.assertIsNone(historical.location)
        self.assertEqual({o.location for o in await self.objects() if o.type == 'building'}, {None, 'A', 'B'})

    async def test_reimport_and_display_name(self):
        entries = services.normalize_row(row())
        first = await services.import_service_connections(entries, self.session)
        second = await services.import_service_connections(entries, self.session)
        self.assertEqual(first['created'], 1)
        self.assertEqual(second['created'], 0)
        self.assertEqual(second['unchanged'], 1)
        self.assertEqual(len(await self.objects()), 2)
        connections = await self.connections()
        self.assertEqual(len(connections), 1)
        self.assertEqual(connections[0].connection_notes, ['PVA: 30 kVA'])
        building = next(o for o in await self.objects() if o.type == 'building')
        self.assertEqual((building.name, building.friendly_name), ('Chäppelimatt5', 'Chäppelimatt 5'))

    async def test_edit(self):
        entries = services.normalize_row(row())
        await services.import_service_connections(entries, self.session)
        entries[0]['source_outgoing'] = ['5']
        entries[0]['connection_notes'] = []
        result = await services.import_service_connections(entries, self.session)
        self.assertEqual(result['updated'], 1)
        connections = await self.connections()
        self.assertEqual(len(connections), 1)
        self.assertEqual(connections[0].source_outgoing, ['5'])
        self.assertEqual(connections[0].connection_notes, [])

    async def test_insurance_numbers_persist_without_becoming_identity(self):
        data = list(row('Haus 1\n\nHaus 1\nHaus 2', 'Scheune\nPVA: 25 kVA'))
        data[2] = '0001A\n\n0001\n0001'
        entries = services.normalize_row(data)
        for _ in range(2):
            result = await services.import_service_connections(entries, self.session)
        buildings = {o.name: o for o in await self.objects() if o.type == 'building'}
        self.assertEqual(buildings['Haus1'].insurance_numbers, ['0001A', '0001'])
        self.assertEqual(buildings['Haus2'].insurance_numbers, ['0001'])
        self.assertIsNone(buildings['Haus1'].ckw_id)
        self.assertEqual(result['objects']['unchanged'], 2)
        self.assertEqual(len(await self.connections()), 2)
        notes = {c.building_id: c.connection_notes for c in await self.connections()}
        self.assertEqual(notes[buildings['Haus1'].id], ['Scheune', 'PVA: 25 kVA'])
        # A partial import must not erase other insured structures at the address.
        entry = {**entries[0], 'insurance_numbers': ['0002']}
        await services.import_service_connections([entry], self.session)
        await services.import_service_connections([{**entry, 'insurance_numbers': None}], self.session)
        self.assertEqual(buildings['Haus1'].insurance_numbers, ['0001A', '0001', '0002'])

    async def test_invalid_positions_skip_before_writing(self):
        data = list(row('A\n\nB', None))
        data[2] = '01\n02\n03'
        result = await services.import_service_connections(services.normalize_row(data), self.session)
        self.assertEqual(result['skipped'], 2)
        self.assertEqual(len(await self.objects()), 1)
        self.assertEqual(await self.connections(), [])

    async def test_singular_insurance_json_and_missing_reference(self):
        entry = services.normalize_row(row())[0]
        entry.pop('insurance_numbers')
        entry['insurance_number'] = '0007B'
        await services.import_service_connections([entry], self.session)
        building = next(o for o in await self.objects() if o.type == 'building')
        self.assertEqual(building.insurance_numbers, ['0007B'])
        result = await services.import_service_connections([
            {**entry, 'insurance_number': '0008', 'source_name': 'Missing'}
        ], self.session)
        self.assertEqual(result['skipped'], 1)
        self.assertEqual(building.insurance_numbers, ['0007B'])

    async def test_municipalities(self):
        for town in ('A', 'B'):
            await services.import_service_connections(services.normalize_row(row(municipality=town)), self.session)
        self.assertEqual(len(await self.objects()), 3)
        self.assertEqual(len(await self.connections()), 2)

    async def test_gps_matches_excel_and_asset_id(self):
        await services.import_service_connections(services.normalize_row(row()), self.session)
        payload = GPSImportData(points=[dict(name='Chäppelimatt 5', lon=8, lat=47, ckw_id='123', type='building', location='A')])
        await services.import_gps_objects(payload, self.session)
        await services.import_gps_objects(payload, self.session)
        payload.points[0].name = 'Neuer Name 5'
        await services.import_gps_objects(payload, self.session)
        objects = await self.objects()
        self.assertEqual(len(objects), 2)
        building = next(o for o in objects if o.ckw_id == '123')
        self.assertIsNotNone(building.geom)
        self.assertEqual(building.type, 'building')
        self.assertEqual(building.name, 'NeuerName5')

    async def test_duplicate_in_batch_last_notes_win(self):
        entry = services.normalize_row(row())[0]
        result = await services.import_service_connections([entry, {**entry, 'connection_notes': ['Last']}], self.session)
        self.assertEqual(result['created'], 1)
        self.assertTrue(result['warnings'])
        self.assertEqual((await self.connections())[0].connection_notes, ['Last'])

    async def test_multiple_connections_reimport_and_ambiguity(self):
        entry = services.normalize_row(row())[0]
        batch = [entry, {**entry, 'source_outgoing': ['5']}]
        await services.import_service_connections(batch, self.session)
        result = await services.import_service_connections(batch, self.session)
        self.assertEqual(result['created'], 0)
        self.assertEqual(len(await self.connections()), 2)
        result = await services.import_service_connections([{**entry, 'source_outgoing': ['6']}], self.session)
        self.assertEqual(result['skipped'], 1)
        self.assertTrue(result['warnings'])
        self.assertEqual(len(await self.connections()), 2)

    async def test_missing_references_reported(self):
        entry = services.normalize_row(row())[0]
        entry.update(source_name='Missing', unswitched_terminal='Missing DP', first_disconnect_point='Missing DB')
        result = await services.import_service_connections([entry], self.session)
        self.assertEqual(result['skipped'], 1)
        self.assertEqual({w['field'] for w in result['warnings']}, {'transformer', 'disconnect_point', 'distribution_box'})
        self.assertEqual(await self.connections(), [])

    async def test_repeated_references_keep_all_outgoing_positions(self):
        self.session.add(Object(name='Box1', type='distribution_box'))
        await self.session.commit()
        entry = services.normalize_row(row())[0]
        entry.update(source_name='Transformer A\nTransformer A',
                     first_disconnect_point='Box 1\nBox 1', source_outgoing=['4', '5'])
        result = await services.import_service_connections([entry], self.session)
        self.assertEqual(result['created'], 1)
        self.assertEqual(result['warnings'], [])
        self.assertEqual((await self.connections())[0].source_outgoing, ['4', '5'])

    async def test_distinct_multiline_box_references_stay_unresolved(self):
        self.session.add_all([Object(name=name, type='distribution_box') for name in ('Box1', 'Box2')])
        await self.session.commit()
        entry = services.normalize_row(row())[0]
        entry['first_disconnect_point'] = 'Box 1\nBox 2'
        result = await services.import_service_connections([entry], self.session)
        self.assertEqual(result['skipped'], 1)
        self.assertEqual(result['warnings'][0]['reason'], 'multiple_distinct_references')
        self.assertEqual(await self.connections(), [])

    async def test_reference_abbreviation_case_and_municipality(self):
        box = Object(name='Eggstrasse11', location='A', type='distribution_box')
        self.session.add_all([box, Object(name='Eggstrasse11', location='B', type='distribution_box')])
        await self.session.commit()
        entry = services.normalize_row(row())[0]
        entry['first_disconnect_point'] = 'EGGSTR. 11'
        result = await services.import_service_connections([entry], self.session)
        self.assertEqual(result['created'], 1)
        self.assertEqual((await self.connections())[0].distribution_box_id, box.id)
        self.assertEqual(result['warnings'][0]['reason'], 'reference_alias_matched')
        self.assertEqual(box.name, 'Eggstrasse11')

    async def test_alias_collision_and_different_house_number_are_not_guessed(self):
        self.session.add_all([Object(name=n, type='distribution_box', location='A')
                              for n in ('Eggstrasse11', 'EGGSTRASSE11')])
        await self.session.commit()
        entry = services.normalize_row(row())[0]
        for name, reason in [('Eggstr. 11', 'ambiguous_object_reference'), ('Eggstr. 12', 'object_not_found')]:
            result = await services.import_service_connections([{**entry, 'first_disconnect_point': name}], self.session)
            self.assertEqual(result['skipped'], 1)
            self.assertEqual(result['warnings'][0]['reason'], reason)

    async def test_reference_alias_index_follows_rename(self):
        box = Object(name='Eggstr. A', ckw_id='box-asset', type='distribution_box')
        self.session.add(box)
        await self.session.flush()
        objects = ObjectIndex([box])
        await upsert_object(self.session, objects, name='Feldstr. B', ckw_id='box-asset')
        self.assertIsNone(find_reference_object(objects, 'Eggstrasse A')[0])
        self.assertIs(find_reference_object(objects, 'Feldstrasse B')[0], box)

    async def test_identical_connection_duplicates_update_without_deleting_ids(self):
        entry = services.normalize_row(row())[0]
        await services.import_service_connections([entry], self.session)
        original = (await self.connections())[0]
        copy = ServiceConnection(building_id=original.building_id, transformer_id=original.transformer_id,
            source_outgoing=original.source_outgoing, disconnect_point_outgoing=original.disconnect_point_outgoing,
            connection_notes=original.connection_notes)
        self.session.add(copy)
        await self.session.commit()
        ids = {original.id, copy.id}
        entry['connection_notes'] = ['Corrected']
        result = await services.import_service_connections([entry], self.session)
        self.assertEqual(result['updated'], 1)
        connections = await self.connections()
        self.assertEqual({c.id for c in connections}, ids)
        self.assertTrue(all(c.connection_notes == ['Corrected'] for c in connections))
        self.assertEqual(result['warnings'][0]['reason'], 'duplicate_connections_synchronized')
        self.assertEqual(set(result['warnings'][0]['connection_ids']), ids)
        result = await services.import_service_connections([entry], self.session)
        self.assertEqual(result['unchanged'], 1)

    async def test_duplicate_connections_with_different_notes_stay_ambiguous(self):
        entry = services.normalize_row(row())[0]
        await services.import_service_connections([entry], self.session)
        original = (await self.connections())[0]
        copy = ServiceConnection(building_id=original.building_id, transformer_id=original.transformer_id,
            source_outgoing=original.source_outgoing, disconnect_point_outgoing=original.disconnect_point_outgoing,
            connection_notes=['Different note'])
        self.session.add(copy)
        await self.session.commit()
        result = await services.import_service_connections([{**entry, 'connection_notes': ['New']}], self.session)
        self.assertEqual(result['skipped'], 1)
        self.assertEqual(original.connection_notes, ['PVA: 30 kVA'])
        self.assertEqual(copy.connection_notes, ['Different note'])

    async def test_partial_connection_payload_does_not_merge_different_paths(self):
        entry = services.normalize_row(row())[0]
        await services.import_service_connections([entry], self.session)
        original = (await self.connections())[0]
        self.session.add(ServiceConnection(building_id=original.building_id, transformer_id=original.transformer_id,
            source_outgoing=['99'], disconnect_point_outgoing=original.disconnect_point_outgoing,
            connection_notes=original.connection_notes))
        await self.session.commit()
        entry.pop('source_outgoing')
        result = await services.import_service_connections([entry], self.session)
        self.assertEqual(result['skipped'], 1)

    async def test_none_preserves_arrays(self):
        entry = services.normalize_row(row())[0]
        await services.import_service_connections([entry], self.session)
        await services.import_service_connections([{**entry, 'source_outgoing': None, 'connection_notes': None}], self.session)
        connection = (await self.connections())[0]
        self.assertEqual(connection.source_outgoing, ['4'])
        self.assertEqual(connection.connection_notes, ['PVA: 30 kVA'])

    async def test_null_items_and_normalized_arrays(self):
        entry = services.normalize_row(row())[0]
        entry['source_outgoing'] = [' 8 ', None, '9']
        await services.import_service_connections([entry], self.session)
        entry['source_outgoing'] = ['8', '9']
        result = await services.import_service_connections([entry], self.session)
        self.assertEqual(result['unchanged'], 1)
        self.assertEqual((await self.connections())[0].source_outgoing, ['8', '9'])

    async def test_legacy_blank_location_and_spaced_name(self):
        self.session.add(Object(name='Chäppelimatt 5', type='building', ckw_id='asset'))
        await self.session.commit()
        await services.import_service_connections(services.normalize_row(row()), self.session)
        building = next(o for o in await self.objects() if o.type == 'building')
        self.assertEqual((building.name, building.location, building.ckw_id), ('Chäppelimatt5', 'A', 'asset'))
        self.assertEqual(len(await self.objects()), 2)

    async def test_ambiguous_names_and_duplicate_assets(self):
        self.session.add_all([Object(name='Haus1', type='building', location=town, ckw_id='same') for town in ('A', 'B')])
        await self.session.commit()
        for point in [dict(name='Haus 1', lat=47, lon=8), dict(name='Renamed', ckw_id='same', lat=47, lon=8)]:
            result = await services.import_gps_objects(GPSImportData(points=[point]), self.session, return_report=True)
            self.assertEqual(result['skipped'], 1)
        self.assertEqual(len(await self.objects()), 3)

    async def test_conflicting_asset_id_not_overwritten(self):
        self.session.add(Object(name='Haus1', type='building', ckw_id='one'))
        await self.session.commit()
        result = await services.import_gps_objects(GPSImportData(points=[dict(name='Haus 1', ckw_id='two', lat=47, lon=8)]), self.session, return_report=True)
        self.assertEqual(result['skipped'], 1)
        self.assertEqual(len(await self.objects()), 2)

    async def test_gps_preserves_missing_attributes_and_friendly_name(self):
        await services.import_service_connections(services.normalize_row(row()), self.session)
        payload = GPSImportData(points=[dict(name='Chäppelimatt5', lon=8, lat=47, ckw_id='123')])
        await services.import_gps_objects(payload, self.session)
        payload.points[0].ckw_id = None
        await services.import_gps_objects(payload, self.session)
        building = next(o for o in await self.objects() if o.type == 'building')
        self.assertEqual((building.ckw_id, building.friendly_name, building.location), ('123', 'Chäppelimatt 5', 'A'))

    async def test_multiple_sources_warning_retains_names(self):
        data = list(row())
        data[7] = 'Transformer A\nTransformer B'
        result = await services.import_service_connections(services.normalize_row(data), self.session)
        warning = next(w for w in result['warnings'] if w['reason'] == 'multiple_sources_single_transformer')
        self.assertEqual(warning['values'], ['Transformer A', 'Transformer B'])
        self.assertEqual(result['created'], 1)

    async def test_positional_notes_persist(self):
        entries = services.normalize_row(row('Haus 1\nHaus 2\nHaus 3', '\n\nLadestation: 11 kVA'))
        await services.import_service_connections(entries, self.session)
        objects = {o.id: o.name for o in await self.objects()}
        self.assertEqual({objects[c.building_id]: c.connection_notes for c in await self.connections()},
                         {'Haus1': [], 'Haus2': [], 'Haus3': ['Ladestation: 11 kVA']})

    async def test_excel_route_reimport(self):
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = 'Modified'
        worksheet.append(['Gemeinde', 'Objekt', 'Assek', 'Bemerkung', 'TK', 'Trennstelle', 'Abgang', 'Speisung', 'Abgang'])
        values = list(row('Haus 1\nHaus 2\nHaus 3', '\n\nLadestation: 11 kVA'))
        values[2] = '0033\n0034\n0035'
        worksheet.append(values)
        data = BytesIO()
        workbook.save(data)
        workbook.close()
        for _ in range(2):
            result = await import_excel(UploadFile(file=BytesIO(data.getvalue()), filename='test.xlsx'), self.session)
        self.assertEqual(result['created'], 0)
        self.assertEqual(result['unchanged'], 3)
        self.assertEqual(result['warnings'], [])
        self.assertEqual(len(await self.connections()), 3)
        self.assertEqual({o.name: o.insurance_numbers for o in await self.objects() if o.type == 'building'},
                         {'Haus1': ['0033'], 'Haus2': ['0034'], 'Haus3': ['0035']})

    async def test_optional_reference_update_preserve_and_clear(self):
        self.session.add_all([Object(name=name, type='distribution_box') for name in ('Box1', 'Box2')])
        await self.session.commit()
        entry = services.normalize_row(row())[0]
        entry['first_disconnect_point'] = 'Box 1'
        await services.import_service_connections([entry], self.session)
        original = (await self.connections())[0].distribution_box_id
        entry['first_disconnect_point'] = None
        await services.import_service_connections([entry], self.session)
        self.assertEqual((await self.connections())[0].distribution_box_id, original)
        entry['first_disconnect_point'] = 'Box 2'
        await services.import_service_connections([entry], self.session)
        self.assertNotEqual((await self.connections())[0].distribution_box_id, original)
        entry['first_disconnect_point'] = ''
        await services.import_service_connections([entry], self.session)
        self.assertIsNone((await self.connections())[0].distribution_box_id)
        self.assertEqual(len(await self.connections()), 1)

    async def test_legacy_idempotence(self):
        data = ImportData(raw_entries=services.normalize_row(row()))
        with patch.object(services, 'AsyncSessionLocal', self.sessions):
            await services.import_to_db(data)
            result = await services.import_to_db(data)
        self.assertEqual(result['created'], 0)
        self.assertEqual(len(await self.objects()), 2)
        self.assertEqual(len(await self.connections()), 1)

    async def test_unexpected_failure_rolls_back(self):
        original = self.session.flush
        calls = 0

        async def failing_flush(*args, **kwargs):
            nonlocal calls
            calls += 1
            if calls == 2:
                raise RuntimeError('test failure')
            return await original(*args, **kwargs)

        with patch.object(self.session, 'flush', side_effect=failing_flush):
            with self.assertRaises(RuntimeError):
                await services.import_service_connections(services.normalize_row(row()), self.session)
        self.assertEqual(len(await self.objects()), 1)
        self.assertEqual(await self.connections(), [])


@unittest.skipUnless(os.getenv('TEST_POSTGRES_URL'), 'Set TEST_POSTGRES_URL for PostGIS integration tests')
class PostgresImportTests(ImportTests):
    """The same regressions on real JSONB, geometry, and PostgreSQL transactions.

    Use a random isolated schema; never read/write the application's data tables.
    The database must already provide PostGIS and the object_type enum.
    """
    async def asyncSetUp(self):
        self.schema = 'import_test_' + uuid.uuid4().hex
        self.engine = create_async_engine(os.environ['TEST_POSTGRES_URL'],
                                          execution_options={'schema_translate_map': {None: self.schema}})
        async with self.engine.begin() as conn:
            await conn.execute(text(f'CREATE SCHEMA "{self.schema}"'))
            await conn.run_sync(lambda c: Object.__table__.create(c))
            await conn.run_sync(lambda c: ServiceConnection.__table__.create(c))
        self.sessions = async_sessionmaker(self.engine, expire_on_commit=False)
        self.session = self.sessions()
        self.session.add(Object(name='TransformerA', friendly_name='Transformer A', type='transformer'))
        await self.session.commit()

    async def asyncTearDown(self):
        await self.session.close()
        async with self.engine.begin() as conn:
            # Only this test's generated schema can be dropped.
            assert self.schema.startswith('import_test_') and len(self.schema) == 44
            await conn.execute(text(f'DROP SCHEMA "{self.schema}" CASCADE'))
        await self.engine.dispose()

    async def test_concurrent_imports(self):
        async def run():
            async with self.sessions() as session:
                return await services.import_service_connections(services.normalize_row(row()), session)
        results = await asyncio.gather(run(), run())
        self.assertEqual(sum(r['created'] for r in results), 1)
        self.assertEqual(len(await self.objects()), 2)
        self.assertEqual(len(await self.connections()), 1)


if __name__ == '__main__':
    unittest.main()
